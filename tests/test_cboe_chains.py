"""Cboe delayed chains: parser (real fixture), planner window, symbol list, storage
lifecycle (parts -> compaction -> 45-day retention), chain features and the slide report."""

from __future__ import annotations

import json
from datetime import UTC, date, datetime, timedelta

import pandas as pd
import pytest
from pptx import Presentation

from omi import calendar as cal
from omi.config import PROJECT_ROOT, load_settings
from omi.features import chains as chain_features
from omi.http import FetchResult
from omi.quality import conform
from omi.report import build_report, fmt_num, fmt_signed
from omi.runner import run_source
from omi.sources.base import NotAvailable, RunContext, SchemaDrift
from omi.sources.cboe_chains import CHAIN, CboeChains, parse_chain, parse_osi
from omi.sources.occ_contract_volume import CONTRACT_VOLUME
from omi.sources.symbol_directory import LISTED
from omi.storage import IngestLog, Storage

S1, S2, S3 = date(2026, 10, 26), date(2026, 10, 27), date(2026, 10, 28)


def evening(d: date, hh: int = 19, mm: int = 30) -> datetime:
    return datetime(d.year, d.month, d.day, hh, mm, tzinfo=cal.ET)


@pytest.fixture
def settings():
    return load_settings(PROJECT_ROOT / "config" / "config.example.yaml")


# ---- parser -----------------------------------------------------------------------------

def test_parse_real_fixture(fixture_bytes):
    out = parse_chain(fixture_bytes("cboe_chains/_VIX_trimmed.json"), S3, "VIX")
    chain = out["chain_eod"].set_index("option")
    assert len(chain) == 3
    row = chain.loc["VIX261118C00010000"]
    assert (row.root, row.expiry, row.cp, row.strike) == ("VIX", date(2026, 11, 18), "C", 10.0)
    assert (row.open_interest, row.volume, row.bid, row.ask) == (4298, 1, 8.35, 8.75)
    assert pd.isna(row.iv)  # Cboe sent 0.0: stored as missing, not zero vol
    assert chain.loc["VIX261118P00010000", "iv"] == pytest.approx(1.0421)
    # last_trade_time is ET wall time: 12:44:59 ET on 29 Oct = 16:44:59 UTC
    assert row.last_trade_at == pd.Timestamp("2026-10-29 16:44:59", tz="UTC")
    under = out["chain_underlying"].iloc[0]
    assert (under.cboe_symbol, under.security_type, under.price) == ("^VIX", "index", 17.2)
    assert under.cboe_timestamp == pd.Timestamp("2026-10-29 20:27:39", tz="UTC")
    conform(CHAIN, out["chain_eod"])  # matches the table schema


def test_parse_osi_handles_long_roots_and_rejects_junk():
    df = parse_osi(pd.Series(["SPXW261118C05800000", "BRKB270115P00450500"]))
    assert df["root"].tolist() == ["SPXW", "BRKB"]
    assert df["strike"].tolist() == [5800.0, 450.5]
    with pytest.raises(SchemaDrift):
        parse_osi(pd.Series(["NOT-AN-OSI"]))


def test_empty_chain_and_html_page_are_not_available():
    with pytest.raises(NotAvailable):
        parse_chain(json.dumps({"data": {"options": []}}).encode(), S3, "X")
    with pytest.raises(NotAvailable):
        parse_chain(b"<?xml version='1.0'?><Error>AccessDenied</Error>", S3, "X")
    with pytest.raises(SchemaDrift):
        parse_chain(json.dumps({"something": 1}).encode(), S3, "X")


# ---- planner ----------------------------------------------------------------------------

def _ctx(settings, started_at, latest=None):
    tables = latest or {}
    return RunContext(settings, started_at.date(),
                      read_latest=lambda n: tables.get(n, pd.DataFrame()),
                      started_at=started_at)


@pytest.mark.parametrize(("started", "expect"), [
    (evening(S3), True),  # same evening after the close: snapshot
    (evening(S3, 15, 0), False),  # before the window (and before the close)
    (evening(S3 + timedelta(days=1), 7, 30), False),  # next-morning retry: skip
])
def test_plan_only_in_the_snapshot_window(settings, started, expect):
    src = CboeChains()
    session = cal.last_completed_session(started)
    reqs = src.plan([session], _ctx(settings, started))
    assert bool(reqs) is expect
    if expect:
        assert {r.as_of for r in reqs} == {S3}
        assert reqs[0].url.endswith("/options/_SPX.json")  # universe first; index gets "_"


def test_symbol_list_order_index_prefix_and_share_class(settings):
    occ = pd.DataFrame([
        {"trade_date": S3, "underlying": u, "product_kind": k, "cust_sides": n,
         "firm_sides": 0, "mm_sides": 0}
        for u, k, n in [("AAPL", "OSTK", 900), ("NDX", "OIND", 500), ("BRKB", "OSTK", 100),
                        ("DEAD", "OSTK", 0)]
    ])
    listings = pd.DataFrame({"as_of": [S3], "symbol": ["BRK.B"], "occ_symbol": ["BRKB"]})
    ctx = _ctx(settings, evening(S3), {"occ_contract_volume": occ,
                                       "ref_listed_symbols": listings})
    pairs = CboeChains().symbols(ctx)
    keys = [p[0] for p in pairs]
    assert keys[:5] == ["SPX", "SPY", "QQQ", "IWM", "VIX"]  # universe first
    assert keys[5:] == ["AAPL", "NDX", "BRKB"]  # then by OCC volume; zero-volume names out
    urls = dict(pairs)
    assert (urls["SPX"], urls["NDX"], urls["BRKB"], urls["AAPL"]) == (
        "_SPX", "_NDX", "BRK.B", "AAPL")


# ---- runner lifecycle -------------------------------------------------------------------

def synth_chain(sym: str, oi_shift: int, n_strikes: int = 5) -> bytes:
    options = []
    for i in range(n_strikes):
        strike = 100 + 5 * i
        for cp in ("C", "P"):
            options.append({
                "option": f"{sym}261218{cp}{strike * 1000:08d}", "bid": 1.0, "ask": 1.2,
                "open_interest": 1000 + 10 * i + oi_shift * (1 if cp == "C" else -1),
                "volume": 50, "iv": 0.3, "delta": 0.5, "gamma": 0.01, "vega": 0.1,
                "theta": -0.02, "theo": 1.1, "last_trade_price": 1.1, "prev_day_close": 1.0,
                "last_trade_time": "2026-10-28T15:59:00",
            })
    return json.dumps({"timestamp": "2026-10-28 20:30:00",
                       "data": {"symbol": sym, "security_type": "stock",
                                "current_price": 110.0, "prev_day_close": 109.0, "iv30": 30.0,
                                "options": options}}).encode()


class ChainHttp:
    def __init__(self, shift: int, missing: set[str] = frozenset()):
        self.shift, self.missing, self.calls = shift, missing, []

    def get(self, url):
        self.calls.append(url)
        sym = url.rsplit("/", 1)[1].removesuffix(".json").lstrip("_")
        if sym in self.missing:
            return FetchResult(url=url, status=403, content=b"<Error/>", headers={},
                               final_url=url, elapsed_s=0, attempts=1, error_kind="http",
                               error="HTTP 403 Forbidden")
        return FetchResult(url=url, status=200, content=synth_chain(sym, self.shift),
                           headers={}, final_url=url, elapsed_s=0, attempts=1)


def run_day(settings, storage, session, shift, missing=frozenset()):
    ctx = RunContext(settings, session, read_latest=storage.read_latest,
                     started_at=evening(session))
    http = ChainHttp(shift, set(missing))
    ingest = IngestLog(storage.meta_dir, f"r{session:%Y%m%d}")
    summary = run_source(CboeChains(), [session], ctx, http=http, storage=storage,
                         ingest=ingest)
    return summary, http


@pytest.fixture
def chain_store(tmp_path, settings):
    settings.sources["cboe_chains"].__pydantic_extra__["scope"] = "universe"
    settings.sources["cboe_chains"].__pydantic_extra__["retention_sessions"] = 2
    settings.sources["cboe_chains"].__pydantic_extra__["keep_raw_sessions"] = 1
    storage = Storage(tmp_path / "data")
    results = [run_day(settings, storage, d, shift, missing={"IWM"} if d == S1 else set())
               for d, shift in ((S1, 0), (S2, 20), (S3, 50))]
    return storage, settings, results


def test_lifecycle_parts_compaction_retention(chain_store):
    storage, _, results = chain_store
    first, _ = results[0]
    assert (first.ok, first.not_available, first.failed) == (4, 1, 0)  # 403 -> not available
    days = sorted(p.name for p in storage.table_dir("chain_eod").iterdir())
    assert days == [f"day={S2}", f"day={S3}"]  # rolling window of 2 sessions
    leftovers = list(storage.table_dir("chain_eod").rglob("part-*.parquet"))
    assert leftovers == []  # all parts compacted into data.parquet
    df = storage.read_table("chain_eod")
    assert len(df) == 2 * 5 * 10  # 2 days x 5 symbols x 10 contracts
    assert not df.duplicated(["as_of", "underlying", "option"]).any()
    raw_days = {p.name[:10] for p in (storage.raw_dir / "cboe_chains").rglob("*.gz")}
    assert raw_days == {f"{S3}"}  # raw JSON kept for 1 session only


def test_rerun_same_evening_skips_stored_symbols(chain_store):
    storage, settings, _ = chain_store
    summary, http = run_day(settings, storage, S3, 50)
    assert summary.skipped == 5 and http.calls == []


def test_chain_features_and_report(chain_store, tmp_path):
    storage, settings, _ = chain_store
    counts = chain_features.build(storage, settings)
    assert counts["chain_underlying_daily"] == 10
    und = pd.read_parquet(storage.features_dir / "chain_underlying_daily.parquet")
    spy = und[(und.underlying == "SPY") & (und.as_of == S3)].iloc[0]
    # calls +30 per strike, puts -30 per strike between S2 and S3: net zero, split visible
    assert spy.d_oi == 0 and spy.call_oi - spy.put_oi == 2 * 5 * 50
    assert spy.prev_as_of == S2
    assert spy.premium_est == pytest.approx(10 * 50 * 1.1 * 100)
    chg = pd.read_parquet(storage.features_dir / "chain_contract_changes.parquet")
    assert len(chg) == 5 * 10 and set(chg.d_oi.abs()) == {30}

    settings.paths.exports_dir = tmp_path / "exports"
    path = build_report(storage, settings)
    prs = Presentation(str(path))
    titles = [s.shapes.title.text for s in prs.slides]
    assert len(titles) == 5 and titles[-1] == "Largest contract-level open-interest changes"


def test_compaction_keeps_newest_row(tmp_path):
    storage = Storage(tmp_path)
    t0 = datetime(2026, 10, 28, 21, tzinfo=UTC)
    raw = parse_chain(synth_chain("SPY", 0), S3, "SPY")["chain_eod"]
    storage.upsert(CHAIN, conform(CHAIN, raw, t0))
    later = raw.assign(open_interest=raw.open_interest + 1)
    storage.upsert(CHAIN, conform(CHAIN, later, t0 + timedelta(hours=1)))
    assert len(list(storage.table_dir("chain_eod").rglob("part-*.parquet"))) == 2
    storage.compact(CHAIN)
    df = storage.read_table("chain_eod")
    assert len(df) == len(raw)
    assert (df.sort_values("option").open_interest.values
            == later.sort_values("option").open_interest.values).all()


def test_ingest_log_compaction(tmp_path):
    log = IngestLog(tmp_path, "run")
    for i in range(5):
        log.add({"source": "s", "report": "r", "as_of_date": S3, "key": str(i), "status": "ok",
                 "bytes": 1, "rows_parsed": 1, "started_at": datetime.now(UTC),
                 "finished_at": datetime.now(UTC)})
        log.flush()
    assert log.compact(max_parts=10) == 0
    assert log.compact(max_parts=3) == 5
    assert len(list(log.dir.glob("part-*.parquet"))) == 1
    assert len(log.ok_keys()) == 5


def test_number_formatting():
    assert fmt_num(1_234_567) == "1.2M" and fmt_num(-12_300) == "-12.3K"
    assert fmt_signed(2500) == "+2.5K" and fmt_signed(0) == "0"


def test_report_without_chains_uses_occ_only(tmp_path, settings):
    storage = Storage(tmp_path / "data")
    rows = [{"trade_date": S3, "underlying": "SPY", "symbol": "SPY", "exchange": "C",
             "product_kind": "OSTK", "expiry": S3, "cust_sides": 60, "firm_sides": 0,
             "mm_sides": 40}]
    storage.upsert(CONTRACT_VOLUME, conform(CONTRACT_VOLUME, pd.DataFrame(rows)))
    storage.upsert(LISTED, conform(LISTED, pd.DataFrame([{
        "as_of": S3, "listing_file": "otherlisted", "symbol": "SPY", "occ_symbol": "SPY",
        "name": "SPDR S&P 500", "listing_exchange": "P", "is_etf": True,
        "is_test_issue": False}])))
    from omi.features.product_class import build

    build(storage, settings)
    settings.paths.exports_dir = tmp_path / "exports"
    prs = Presentation(str(build_report(storage, settings)))
    assert len(prs.slides) == 5
