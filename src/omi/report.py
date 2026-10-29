"""Daily options-positioning slides (``omi report``), built with python-pptx.

Reads only feature files, so the deck always matches the dashboard. Charts are native
PowerPoint charts (editable, restyle-able), tables are native tables. If ``report.template``
points to a .pptx, its layouts and theme are used, so it picks up the corporate look.

Slides:
  1. Title: as-of date, sources, data caveats.
  2. Market activity (OCC, all exchanges): contracts per day by class, stacked, last N
     sessions, plus headline numbers.
  3. Largest open-interest builds by underlying (chains).
  4. Largest open-interest unwinds by underlying (chains).
  5. Largest contract-level open-interest changes (chains).
Chain slides are replaced by a one-line note when no chain snapshots are stored yet.
"""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

import duckdb
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from omi.config import Settings
from omi.storage import Storage

log = logging.getLogger(__name__)

CLASS_LABELS = {"etf": "ETF", "single_stock": "Single stock", "index": "Index"}
CLASS_COLOURS = {"ETF": "2A78D6", "Single stock": "EB6834", "Index": "1BAF7A"}
GREY = RGBColor(0x59, 0x59, 0x59)


class ReportError(Exception):
    """Nothing to report (no features built yet)."""


# ---- formatting -------------------------------------------------------------------------

def fmt_num(x: float | int | None) -> str:
    """12,345 -> '12.3K'; keeps sign; blanks for missing."""
    if x is None or pd.isna(x):
        return ""
    sign = "-" if x < 0 else ""
    x = abs(float(x))
    for div, suffix in ((1e12, "T"), (1e9, "B"), (1e6, "M"), (1e3, "K")):
        if x >= div:
            return f"{sign}{x / div:.1f}{suffix}"
    return f"{sign}{x:,.0f}"


def fmt_signed(x: float | int | None) -> str:
    text = fmt_num(x)
    return f"+{text}" if text and not text.startswith("-") and float(x) != 0 else text


def fmt_pct(x: float | None, signed: bool = True) -> str:
    if x is None or pd.isna(x):
        return ""
    return f"{x:+.1%}" if signed else f"{x:.1%}"


def fmt_money(x: float | None) -> str:
    text = fmt_num(x)
    if not text:
        return ""
    return f"-${text[1:]}" if text.startswith("-") else f"${text}"


# ---- data -------------------------------------------------------------------------------

def _q(sql: str) -> pd.DataFrame:
    with duckdb.connect() as con:
        return con.execute(sql).df()


def _feature(storage: Storage, name: str) -> Path | None:
    path = storage.features_dir / f"{name}.parquet"
    return path if path.exists() else None


def market_activity(storage: Storage, sessions: int) -> pd.DataFrame:
    """trade_date x class -> contracts, customer share, for the last ``sessions`` days."""
    path = _feature(storage, "class_daily")
    if path is None:
        return pd.DataFrame()
    df = _q(f"""
        WITH d AS (
            SELECT trade_date, product_class, sum(contracts) AS contracts,
                   sum(cust_sides) AS cust, sum(cust_sides + firm_sides + mm_sides) AS sides
            FROM read_parquet('{path.as_posix()}')
            WHERE product_class IN ('etf', 'single_stock', 'index')
            GROUP BY trade_date, product_class
        ), days AS (SELECT DISTINCT trade_date FROM d ORDER BY trade_date DESC LIMIT {sessions})
        SELECT d.* FROM d JOIN days USING (trade_date) ORDER BY trade_date, product_class
    """)
    df["class"] = df["product_class"].map(CLASS_LABELS)
    return df


def chain_tables(storage: Storage, as_of: date | None, top: int) -> dict[str, pd.DataFrame]:
    und = _feature(storage, "chain_underlying_daily")
    chg = _feature(storage, "chain_contract_changes")
    if und is None:
        return {}
    day = f"DATE '{as_of}'" if as_of else (
        f"(SELECT max(as_of) FROM read_parquet('{und.as_posix()}'))")
    latest = _q(f"SELECT * FROM read_parquet('{und.as_posix()}') "
                f"WHERE as_of = {day} AND d_oi IS NOT NULL")
    out = {
        "builds": latest[latest.d_oi > 0].nlargest(top, "d_oi"),
        "unwinds": latest[latest.d_oi < 0].nsmallest(top, "d_oi"),
    }
    if chg is not None:
        c = _q(f"SELECT * FROM read_parquet('{chg.as_posix()}') WHERE as_of = {day}")
        if not c.empty:
            c = c.reindex(c["d_oi"].abs().sort_values(ascending=False).index).head(top)
        out["contracts"] = c
    return out


# ---- slide helpers ----------------------------------------------------------------------

def _layout(prs: Presentation, *names: str):
    for name in names:
        for layout in prs.slide_layouts:
            if layout.name.lower() == name.lower():
                return layout
    return prs.slide_layouts[min(5, len(prs.slide_layouts) - 1)]


def _title_slide(prs, slide_title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(_layout(prs, "Title Slide"))
    slide.shapes.title.text = slide_title
    if len(slide.placeholders) > 1:
        slide.placeholders[1].text = subtitle


def _content_slide(prs, slide_title: str, restyle: bool = True):
    slide = prs.slides.add_slide(_layout(prs, "Title Only"))
    title = slide.shapes.title
    title.text = slide_title
    if restyle:  # python-pptx's blank template centres a huge title; use a slim header bar
        title.left, title.top = Inches(0.5), Inches(0.3)
        title.width, title.height = prs.slide_width - Inches(1.0), Inches(0.8)
        para = title.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT
        para.font.size = Pt(28)
        para.font.bold = True
    return slide


def _note(slide, text: str, top: Emu, prs) -> None:
    box = slide.shapes.add_textbox(Inches(0.5), top, prs.slide_width - Inches(1.0), Inches(0.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(11)
    p.font.color.rgb = GREY


def _table(slide, prs, df: pd.DataFrame, headers: list[str], top: Emu,
           widths: list[float] | None = None) -> None:
    rows, cols = len(df) + 1, len(headers)
    width = prs.slide_width - Inches(1.0)
    height = Inches(0.3) * rows
    shape = slide.shapes.add_table(rows, cols, Inches(0.5), top, width, height)
    table = shape.table
    if widths:
        total = sum(widths)
        for i, w in enumerate(widths):
            table.columns[i].width = int(width * w / total)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.text_frame.paragraphs[0].font.size = Pt(11)
        cell.text_frame.paragraphs[0].font.bold = True
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = "" if value is None or (isinstance(value, float) and pd.isna(value)) \
                else str(value)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(10)
            cell.margin_top = cell.margin_bottom = Inches(0.02)


def _class_label(product_class, index_scope) -> str:
    if product_class == "index":
        return f"Index ({index_scope})" if index_scope else "Index"
    return CLASS_LABELS.get(product_class, "") if isinstance(product_class, str) else ""


def _underlying_rows(df: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame({
        "Underlying": df["underlying"],
        "Class": [_class_label(p, s) for p, s in zip(df["product_class"], df["index_scope"],
                                                      strict=True)],
        "Open interest": df["total_oi"].map(fmt_num),
        "Change": df["d_oi"].map(fmt_signed),
        "Change %": df["d_oi_pct"].map(fmt_pct),
        "5-session change": df["d_oi_5"].map(fmt_signed),
        "Put/call OI": df["pc_oi"].map(lambda v: "" if pd.isna(v) else f"{v:.2f}"),
        "Volume / OI": df["vol_to_oi"].map(lambda v: "" if pd.isna(v) else f"{v:.2f}"),
    })


# ---- build ------------------------------------------------------------------------------

def build_report(storage: Storage, settings: Settings, *, as_of: date | None = None,
                 out: Path | None = None) -> Path:
    cfg = settings.report
    activity = market_activity(storage, cfg.sessions)
    chains = chain_tables(storage, as_of, cfg.top_n)
    if activity.empty and not chains:
        raise ReportError("No features built yet; run `omi daily` or `omi features` first.")

    report_day = as_of
    if report_day is None:
        days = []
        if not activity.empty:
            days.append(pd.Timestamp(activity["trade_date"].max()).date())
        if chains.get("builds") is not None and not chains["builds"].empty:
            days.append(pd.Timestamp(chains["builds"]["as_of"].max()).date())
        report_day = max(days) if days else date.today()

    prs = Presentation(str(cfg.template)) if cfg.template else Presentation()
    restyle = cfg.template is None
    if not cfg.template:  # 16:9 instead of python-pptx's default 4:3
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    _title_slide(prs, f"Options market positioning: {report_day:%d %b %Y}",
                 "Sources: OCC (all U.S. exchanges) and Cboe delayed option chains. "
                 "Volumes in contracts; open interest is OCC's start-of-session figure.")

    # 2. market activity
    slide = _content_slide(prs, "Options activity by product class (all U.S. exchanges)",
                           restyle)
    if activity.empty:
        _note(slide, "No OCC contract-volume data stored yet.", Inches(1.5), prs)
    else:
        days = sorted(activity["trade_date"].unique())
        chart_data = CategoryChartData()
        chart_data.categories = [pd.Timestamp(d).strftime("%m/%d") for d in days]
        for label in ("ETF", "Single stock", "Index"):
            part = activity[activity["class"] == label].set_index("trade_date")["contracts"]
            chart_data.add_series(label, [float(part.get(d, 0.0)) / 1e6 for d in days])
        gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.5), Inches(1.4),
                                    Inches(8.3), Inches(5.3), chart_data)
        chart = gf.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.tick_labels.font.size = Pt(10)
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.has_title = True
        chart.value_axis.axis_title.text_frame.text = "contracts (millions)"
        for series in chart.series:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = RGBColor.from_string(CLASS_COLOURS[series.name])
        chart.plots[0].gap_width = 60

        last = activity[activity["trade_date"] == days[-1]]
        total = last["contracts"].sum()
        avg = activity.groupby("trade_date")["contracts"].sum().mean()
        cust = last["cust"].sum() / last["sides"].sum() if last["sides"].sum() else float("nan")
        lines = [f"{pd.Timestamp(days[-1]):%a %d %b}", f"Total: {fmt_num(total)} contracts",
                 f"vs {len(days)}-session avg: {fmt_pct(total / avg - 1)}",
                 f"Customer share: {fmt_pct(cust, signed=False)}", ""]
        ranked = last.sort_values("contracts", ascending=False)
        for label, n in zip(ranked["class"], ranked["contracts"], strict=True):
            lines.append(f"{label}: {fmt_num(n)} ({n / total:.0%})")
        box = slide.shapes.add_textbox(Inches(9.1), Inches(1.6), Inches(3.8), Inches(5.0))
        tf = box.text_frame
        tf.word_wrap = True
        for i, text in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = text
            p.font.size = Pt(16 if i == 0 else 13)
            p.font.bold = i == 0

    # 3-5. open interest
    headers = ["Underlying", "Class", "Open interest", "Change", "Change %", "5-session change",
               "Put/call OI", "Volume / OI"]
    widths = [1.2, 1.4, 1.3, 1.1, 1.0, 1.4, 1.0, 1.0]
    for key, slide_title in (("builds", "Largest open-interest builds by underlying"),
                             ("unwinds", "Largest open-interest unwinds by underlying")):
        slide = _content_slide(prs, slide_title, restyle)
        df = chains.get(key)
        if df is None or df.empty:
            _note(slide, "Needs two days of option-chain snapshots (Cboe delayed chains). "
                         "None stored yet." if not chains else "No changes to show.",
                  Inches(1.5), prs)
            continue
        prev = pd.Timestamp(df["prev_as_of"].iloc[0]).date()
        _table(slide, prs, _underlying_rows(df), headers, Inches(1.3), widths)
        _note(slide, f"Change vs snapshot of {prev:%d %b %Y}. Open interest is OCC's "
                     "start-of-session figure from Cboe's delayed chains.",
              prs.slide_height - Inches(0.7), prs)

    slide = _content_slide(prs, "Largest contract-level open-interest changes", restyle)
    c = chains.get("contracts")
    if c is None or c.empty:
        _note(slide, "Needs two days of option-chain snapshots. None stored yet.",
              Inches(1.5), prs)
    else:
        rows = pd.DataFrame({
            "Underlying": c["underlying"],
            "Expiry": pd.to_datetime(c["expiry"]).dt.strftime("%d %b %y"),
            "Strike": c["strike"].map(lambda v: f"{v:,.2f}".rstrip("0").rstrip(".")),
            "C/P": c["cp"],
            "OI before": c["oi_prev"].map(fmt_num),
            "OI now": c["oi_now"].map(fmt_num),
            "Change": c["d_oi"].map(fmt_signed),
            "~Notional change": c["d_oi_notional"].map(fmt_money),
        })
        _table(slide, prs, rows, list(rows.columns), Inches(1.3),
               [1.2, 1.1, 1.0, 0.6, 1.1, 1.1, 1.1, 1.5])
        _note(slide, "Notional change = OI change x closing mark x 100 (approximate).",
              prs.slide_height - Inches(0.7), prs)

    out = out or (settings.paths.exports_dir / f"omi_report_{report_day:%Y%m%d}.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out
